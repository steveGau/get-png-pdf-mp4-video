"""
Batch Image Cropper
Browse a folder, draw a crop rectangle on a preview image, crop all images,
and export images to an Excel file (cropImages.xlsx).
"""

import os
import tempfile
import tkinter as tk
from tkinter import ttk, filedialog, messagebox

from PIL import Image, ImageTk

try:
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
    from openpyxl.drawing.xdr import XDRPositiveSize2D
    from openpyxl.utils.units import pixels_to_EMU
    from openpyxl.utils import get_column_letter

    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".bmp", ".gif", ".webp", ".tif", ".tiff"}


class BatchImageCropper:
    def __init__(self, root):
        self.root = root
        self.root.title("Batch Image Cropper & Excel Export")
        self.root.geometry("1100x820")
        self.root.minsize(900, 680)
        self.root.configure(bg="#1e1e2e")

        self.folder_path = None
        self.image_paths = []
        self.preview_index = 0
        self.original_image = None
        self.display_photo = None
        self.scale = 1.0
        self.offset_x = 0
        self.offset_y = 0

        self.drag_start = None
        self.crop_rect_id = None
        self.crop_box_image = None  # (x1, y1, x2, y2) in original image pixels

        self._build_ui()

    def _build_ui(self):
        bg = "#1e1e2e"
        fg = "#cdd6f4"
        btn_bg = "#313244"

        top = ttk.Frame(self.root, padding=10)
        top.pack(fill=tk.X)

        ttk.Button(top, text="Browse Folder", command=self.browse_folder).pack(side=tk.LEFT, padx=5)
        self.folder_label = tk.Label(
            top, text="No folder selected", bg=bg, fg=fg, anchor=tk.W
        )
        self.folder_label.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=10)

        nav = ttk.Frame(self.root, padding=(10, 0))
        nav.pack(fill=tk.X)

        ttk.Button(nav, text="◀ Prev", command=self.prev_image).pack(side=tk.LEFT, padx=5)
        ttk.Button(nav, text="Next ▶", command=self.next_image).pack(side=tk.LEFT, padx=5)
        self.image_info_label = tk.Label(nav, text="", bg=bg, fg=fg)
        self.image_info_label.pack(side=tk.LEFT, padx=10)

        ttk.Button(nav, text="Clear Crop Box", command=self.clear_crop_box).pack(side=tk.LEFT, padx=5)

        excel_frame = ttk.LabelFrame(self.root, text="Export Images to Excel (cropImages.xlsx)", padding=10)
        excel_frame.pack(fill=tk.X, padx=10, pady=(5, 0))

        row1 = ttk.Frame(excel_frame)
        row1.pack(fill=tk.X, pady=2)
        ttk.Label(row1, text="Sort by:").pack(side=tk.LEFT, padx=(0, 5))
        self.sort_var = tk.StringVar(value="file name")
        ttk.Combobox(
            row1,
            textvariable=self.sort_var,
            values=["file name", "modify time"],
            state="readonly",
            width=14,
        ).pack(side=tk.LEFT, padx=5)

        ttk.Label(row1, text="Rows:").pack(side=tk.LEFT, padx=(15, 5))
        self.excel_rows_var = tk.StringVar(value="5")
        ttk.Entry(row1, textvariable=self.excel_rows_var, width=6).pack(side=tk.LEFT, padx=2)

        ttk.Label(row1, text="Columns:").pack(side=tk.LEFT, padx=(10, 5))
        self.excel_cols_var = tk.StringVar(value="4")
        ttk.Entry(row1, textvariable=self.excel_cols_var, width=6).pack(side=tk.LEFT, padx=2)

        row2 = ttk.Frame(excel_frame)
        row2.pack(fill=tk.X, pady=2)
        ttk.Label(row2, text="Row gap (px):").pack(side=tk.LEFT, padx=(0, 5))
        self.row_gap_var = tk.StringVar(value="10")
        ttk.Entry(row2, textvariable=self.row_gap_var, width=6).pack(side=tk.LEFT, padx=2)

        ttk.Label(row2, text="Column gap (px):").pack(side=tk.LEFT, padx=(15, 5))
        self.col_gap_var = tk.StringVar(value="10")
        ttk.Entry(row2, textvariable=self.col_gap_var, width=6).pack(side=tk.LEFT, padx=2)

        ttk.Button(row2, text="Export to Excel", command=self.export_images_to_excel).pack(
            side=tk.LEFT, padx=(20, 5)
        )
        tk.Label(
            row2,
            text="Images fill row 1 left→right, then row 2, … (extra rows added if needed)",
            bg="#1e1e2e",
            fg="#a6adc8",
        ).pack(side=tk.LEFT, padx=10)

        hint = tk.Label(
            self.root,
            text="Drag mouse on the image to draw the crop rectangle.",
            bg=bg,
            fg="#a6adc8",
        )
        hint.pack(pady=(5, 0))

        canvas_frame = ttk.Frame(self.root, padding=10)
        canvas_frame.pack(fill=tk.BOTH, expand=True)

        self.canvas = tk.Canvas(
            canvas_frame,
            bg="#11111b",
            highlightthickness=1,
            highlightbackground="#45475a",
            cursor="crosshair",
        )
        self.canvas.pack(fill=tk.BOTH, expand=True)

        self.canvas.bind("<ButtonPress-1>", self.on_press)
        self.canvas.bind("<B1-Motion>", self.on_drag)
        self.canvas.bind("<ButtonRelease-1>", self.on_release)
        self.root.bind("<Configure>", self._on_resize)

        bottom = ttk.Frame(self.root, padding=10)
        bottom.pack(fill=tk.X)

        ttk.Button(
            bottom,
            text="Crop All Images",
            command=self.crop_all_images,
        ).pack(side=tk.LEFT, padx=5)

        self.status_var = tk.StringVar(value="Select a folder containing images.")
        tk.Label(bottom, textvariable=self.status_var, bg=bg, fg=fg, anchor=tk.W).pack(
            side=tk.LEFT, fill=tk.X, expand=True, padx=10
        )

    def browse_folder(self):
        folder = filedialog.askdirectory(title="Select Image Folder")
        if not folder:
            return

        paths = []
        for name in sorted(os.listdir(folder)):
            ext = os.path.splitext(name)[1].lower()
            if ext in IMAGE_EXTENSIONS:
                paths.append(os.path.join(folder, name))

        if not paths:
            messagebox.showwarning("No Images", "No image files found in the selected folder.")
            return

        self.folder_path = folder
        self.image_paths = paths
        self.preview_index = 0
        self.crop_box_image = None
        self.folder_label.config(text=folder)
        self.status_var.set(f"Loaded {len(paths)} image(s). Draw a crop rectangle, then click Crop All Images.")
        self.load_preview_image()

    def load_preview_image(self):
        if not self.image_paths:
            return

        path = self.image_paths[self.preview_index]
        try:
            self.original_image = Image.open(path).convert("RGB")
        except OSError as e:
            messagebox.showerror("Error", f"Could not open image:\n{path}\n\n{e}")
            return

        name = os.path.basename(path)
        size = self.original_image.size
        self.image_info_label.config(
            text=f"Preview: {name}  ({self.preview_index + 1}/{len(self.image_paths)})  {size[0]}x{size[1]} px"
        )
        self._render_preview()

    def _on_resize(self, _event=None):
        if self.original_image is not None:
            self._render_preview()

    def _render_preview(self):
        if self.original_image is None:
            return

        self.canvas.update_idletasks()
        cw = max(self.canvas.winfo_width(), 100)
        ch = max(self.canvas.winfo_height(), 100)
        iw, ih = self.original_image.size

        self.scale = min(cw / iw, ch / ih, 1.0)
        dw = max(1, int(iw * self.scale))
        dh = max(1, int(ih * self.scale))
        self.offset_x = (cw - dw) // 2
        self.offset_y = (ch - dh) // 2

        resized = self.original_image.resize((dw, dh), Image.Resampling.LANCZOS)
        self.display_photo = ImageTk.PhotoImage(resized)

        self.canvas.delete("all")
        self.canvas.create_image(self.offset_x, self.offset_y, anchor=tk.NW, image=self.display_photo)

        if self.crop_box_image:
            self._draw_crop_rect_from_image_box()

    def _canvas_to_image_point(self, cx, cy):
        ix = int((cx - self.offset_x) / self.scale)
        iy = int((cy - self.offset_y) / self.scale)
        iw, ih = self.original_image.size
        ix = max(0, min(ix, iw - 1))
        iy = max(0, min(iy, ih - 1))
        return ix, iy

    def _image_box_to_canvas_rect(self, box):
        x1, y1, x2, y2 = box
        return (
            self.offset_x + x1 * self.scale,
            self.offset_y + y1 * self.scale,
            self.offset_x + x2 * self.scale,
            self.offset_y + y2 * self.scale,
        )

    def _draw_crop_rect_from_image_box(self):
        if not self.crop_box_image:
            return
        if self.crop_rect_id:
            self.canvas.delete(self.crop_rect_id)
        x1, y1, x2, y2 = self._image_box_to_canvas_rect(self.crop_box_image)
        self.crop_rect_id = self.canvas.create_rectangle(
            x1, y1, x2, y2, outline="#89b4fa", width=2, dash=(4, 2)
        )

    def on_press(self, event):
        if self.original_image is None:
            return
        self.drag_start = (event.x, event.y)
        if self.crop_rect_id:
            self.canvas.delete(self.crop_rect_id)
            self.crop_rect_id = None

    def on_drag(self, event):
        if self.drag_start is None or self.original_image is None:
            return
        if self.crop_rect_id:
            self.canvas.delete(self.crop_rect_id)
        self.crop_rect_id = self.canvas.create_rectangle(
            self.drag_start[0],
            self.drag_start[1],
            event.x,
            event.y,
            outline="#89b4fa",
            width=2,
            dash=(4, 2),
        )

    def on_release(self, event):
        if self.drag_start is None or self.original_image is None:
            return

        x1c, y1c = self.drag_start
        x2c, y2c = event.x, event.y
        self.drag_start = None

        ix1, iy1 = self._canvas_to_image_point(x1c, y1c)
        ix2, iy2 = self._canvas_to_image_point(x2c, y2c)

        left, top = min(ix1, ix2), min(iy1, iy2)
        right, bottom = max(ix1, ix2), max(iy1, iy2)
        iw, ih = self.original_image.size

        if right - left < 2 or bottom - top < 2:
            self.crop_box_image = None
            if self.crop_rect_id:
                self.canvas.delete(self.crop_rect_id)
                self.crop_rect_id = None
            self.status_var.set("Crop box too small. Drag again.")
            return

        right = min(right, iw)
        bottom = min(bottom, ih)
        self.crop_box_image = (left, top, right, bottom)
        self._draw_crop_rect_from_image_box()
        self.status_var.set(
            f"Crop box: x={left}, y={top}, w={right - left}, h={bottom - top} (applied to all images on crop)"
        )

    def clear_crop_box(self):
        self.crop_box_image = None
        if self.crop_rect_id:
            self.canvas.delete(self.crop_rect_id)
            self.crop_rect_id = None
        self.status_var.set("Crop box cleared.")

    def prev_image(self):
        if not self.image_paths:
            return
        self.preview_index = (self.preview_index - 1) % len(self.image_paths)
        self.load_preview_image()

    def next_image(self):
        if not self.image_paths:
            return
        self.preview_index = (self.preview_index + 1) % len(self.image_paths)
        self.load_preview_image()

    def crop_all_images(self):
        if not self.folder_path or not self.image_paths:
            messagebox.showwarning("Warning", "Please browse and select an image folder first.")
            return
        if not self.crop_box_image:
            messagebox.showwarning("Warning", "Please draw a crop rectangle on the preview image first.")
            return

        out_dir = os.path.join(self.folder_path, "cropImages")
        os.makedirs(out_dir, exist_ok=True)

        x1, y1, x2, y2 = self.crop_box_image
        ok_count = 0
        fail_count = 0

        for path in self.image_paths:
            name = os.path.basename(path)
            out_path = os.path.join(out_dir, name)
            try:
                with Image.open(path) as img:
                    cropped = img.crop((x1, y1, x2, y2))
                    ext = os.path.splitext(name)[1].lower()
                    if ext in (".jpg", ".jpeg"):
                        cropped.save(out_path, quality=95)
                    else:
                        cropped.save(out_path)
                ok_count += 1
            except OSError:
                fail_count += 1

        msg = f"Cropped {ok_count} image(s).\nSaved to:\n{out_dir}"
        if fail_count:
            msg += f"\n\nFailed: {fail_count} image(s)."
        self.status_var.set(f"Done: {ok_count} cropped, {fail_count} failed.")
        messagebox.showinfo("Crop Complete", msg)

    def _sorted_image_paths(self):
        """Return image paths sorted by combo box selection."""
        paths = list(self.image_paths)
        if self.sort_var.get() == "modify time":
            paths.sort(key=os.path.getmtime)
        else:
            paths.sort(key=lambda p: os.path.basename(p).lower())
        return paths

    def _parse_positive_int(self, value, field_name, minimum=1):
        try:
            n = int(value.strip())
        except ValueError:
            messagebox.showwarning("Invalid Input", f"{field_name} must be a whole number.")
            return None
        if n < minimum:
            messagebox.showwarning("Invalid Input", f"{field_name} must be >= {minimum}.")
            return None
        return n

    def _parse_non_negative_int(self, value, field_name):
        try:
            n = int(value.strip())
        except ValueError:
            messagebox.showwarning("Invalid Input", f"{field_name} must be a whole number.")
            return None
        if n < 0:
            messagebox.showwarning("Invalid Input", f"{field_name} must be >= 0.")
            return None
        return n

    def _image_path_for_excel(self, path):
        """openpyxl needs a path; convert GIF/WebP to temporary PNG if required."""
        ext = os.path.splitext(path)[1].lower()
        if ext in (".png", ".jpg", ".jpeg", ".bmp"):
            return path, False
        img = Image.open(path)
        if img.mode not in ("RGB", "RGBA"):
            img = img.convert("RGB")
        tmp = tempfile.NamedTemporaryFile(suffix=".png", delete=False)
        tmp.close()
        img.save(tmp.name, "PNG")
        return tmp.name, True

    def export_images_to_excel(self):
        if not HAS_OPENPYXL:
            messagebox.showerror(
                "Missing Package",
                "openpyxl is required for Excel export.\n\nInstall with:\npip install openpyxl",
            )
            return
        if not self.folder_path or not self.image_paths:
            messagebox.showwarning("Warning", "Please browse and select an image folder first.")
            return

        cols = self._parse_positive_int(self.excel_cols_var.get(), "Columns")
        if cols is None:
            return
        rows = self._parse_positive_int(self.excel_rows_var.get(), "Rows")
        if rows is None:
            return
        row_gap = self._parse_non_negative_int(self.row_gap_var.get(), "Row gap")
        if row_gap is None:
            return
        col_gap = self._parse_non_negative_int(self.col_gap_var.get(), "Column gap")
        if col_gap is None:
            return

        paths = self._sorted_image_paths()
        if not paths:
            messagebox.showwarning("No Images", "No images to export.")
            return

        self.status_var.set("Preparing images for Excel...")
        self.root.update()

        temp_files = []
        try:
            sizes = []
            excel_image_paths = []
            for path in paths:
                xl_path, is_temp = self._image_path_for_excel(path)
                if is_temp:
                    temp_files.append(xl_path)
                excel_image_paths.append(xl_path)
                with Image.open(xl_path) as img:
                    sizes.append(img.size)

            max_w = max(w for w, _ in sizes)
            max_h = max(h for _, h in sizes)

            wb = Workbook()
            ws = wb.active
            ws.title = "Images"

            margin = 5
            for idx, (src_path, (iw, ih)) in enumerate(zip(excel_image_paths, sizes)):
                grid_row = idx // cols
                grid_col = idx % cols

                x_px = margin + grid_col * (max_w + col_gap)
                y_px = margin + grid_row * (max_h + row_gap)

                xl_img = XLImage(src_path)
                xl_img.width = iw
                xl_img.height = ih
                xl_img.anchor = OneCellAnchor(
                    _from=AnchorMarker(
                        col=0,
                        colOff=pixels_to_EMU(x_px),
                        row=0,
                        rowOff=pixels_to_EMU(y_px),
                    ),
                    ext=XDRPositiveSize2D(
                        pixels_to_EMU(iw),
                        pixels_to_EMU(ih),
                    ),
                )
                ws.add_image(xl_img)

            total_rows = max(rows, (len(paths) + cols - 1) // cols)
            sheet_w = margin * 2 + cols * max_w + max(0, cols - 1) * col_gap
            sheet_h = margin * 2 + total_rows * max_h + max(0, total_rows - 1) * row_gap

            for r in range(1, total_rows + 2):
                ws.row_dimensions[r].height = (max_h + row_gap) * 0.75
            for c in range(1, cols + 2):
                ws.column_dimensions[get_column_letter(c)].width = max(8, (max_w + col_gap) / 7)

            out_path = os.path.join(self.folder_path, "cropImages.xlsx")
            wb.save(out_path)

            self.status_var.set(f"Excel saved: {out_path}")
            messagebox.showinfo(
                "Excel Export Complete",
                f"Exported {len(paths)} image(s) to:\n{out_path}\n\n"
                f"Sort: {self.sort_var.get()}\n"
                f"Grid: {cols} column(s) per row, {total_rows} row(s)\n"
                f"Row gap: {row_gap} px, Column gap: {col_gap} px\n"
                f"Layout size: ~{sheet_w} x {sheet_h} px",
            )
        except OSError as e:
            messagebox.showerror("Error", f"Failed to create Excel file:\n{e}")
            self.status_var.set("Excel export failed.")
        finally:
            for tmp in temp_files:
                try:
                    os.remove(tmp)
                except OSError:
                    pass


def main():
    root = tk.Tk()
    style = ttk.Style()
    try:
        style.theme_use("clam")
    except tk.TclError:
        pass
    BatchImageCropper(root)
    root.mainloop()


if __name__ == "__main__":
    main()
